"""End-to-end tests: PDF bytes in -> validated XLSX out.

Covers the full pipeline with the real open-source engine (no network,
no heavy models) and with the OpenAI-compatible engine driven by a fake
client.
"""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from pdf2xlsx.engines.openai_compat import OpenAICompatibleEngine
from pdf2xlsx.pipeline import convert_pdf

from conftest import (
    BORDERLESS_TABLE,
    DEFAULT_TABLE,
    HAS_PERSIAN_FONT,
    TableSpec,
    TitleSpec,
    build_pdf,
)
from test_openai_engine import JSON_TEXT, make_client


def _workbook(outcome):
    return load_workbook(filename=BytesIO(outcome.xlsx_bytes))


def test_e2e_opensource_multi_layout_document(make_settings, tmp_path):
    """PDF with multiple pages and different table structures -> XLSX."""
    narrow = TableSpec(
        rows=[["K", "V"], ["a", "1"], ["b", ""]],
        origin=(60.0, 420.0), cell_w=130.0, cell_h=22.0,
    )
    pdf_path = build_pdf(
        tmp_path / "mixed.pdf",
        [
            [TitleSpec("Report"), DEFAULT_TABLE],
            [narrow, BORDERLESS_TABLE],
        ],
    )
    outcome = convert_pdf(pdf_path.read_bytes(), "opensource", make_settings(),
                          filename="mixed.pdf")

    assert outcome.result.page_count == 2
    assert outcome.result.table_count == 3
    assert outcome.validation.ok
    assert outcome.xlsx_filename == "mixed.xlsx"

    wb = _workbook(outcome)
    assert wb.sheetnames == ["P1T1", "P2T1", "P2T2"]

    ruled = wb["P1T1"]
    assert ruled.max_column == 4
    assert [ruled.cell(row=1, column=c).value for c in range(1, 5)] == \
        ["ID", "Name", "Qty", "Price"]
    assert ruled.cell(row=4, column=2).value is None  # empty cell preserved
    assert ruled.cell(row=4, column=4).value == "1.25"

    # page 2 reading order: borderless (y=100) first, narrow ruled (y=420) second
    borderless = wb["P2T1"]
    assert borderless.max_column == 3  # different layout than page 1
    assert borderless.cell(row=2, column=1).value == "A-1"
    assert borderless.cell(row=4, column=3).value == "8"

    narrow_ws = wb["P2T2"]
    assert narrow_ws.max_column == 2
    assert narrow_ws.cell(row=3, column=2).value is None  # empty cell


def test_e2e_headers_detected_and_written(make_settings, tmp_path):
    pdf_path = build_pdf(tmp_path / "hdr.pdf", [[DEFAULT_TABLE]])
    outcome = convert_pdf(pdf_path, "opensource", make_settings())
    _, table = next(outcome.result.iter_tables())
    assert [c.header for c in table.columns] == ["ID", "Name", "Qty", "Price"]
    assert table.rows[0].is_header is True


@pytest.mark.skipif(not HAS_PERSIAN_FONT, reason="no Persian-capable font on this system")
def test_e2e_persian_pdf_through_opensource(make_settings, tmp_path):
    """A real Persian table in a PDF -> Persian XLSX cells."""
    persian_table = TableSpec(
        rows=[
            ["ردیف", "شرح"],
            ["۱", "پرداخت"],
            ["۲", "بستانکار"],
        ],
        origin=(80.0, 80.0),
        cell_w=150.0,
        cell_h=26.0,
    )
    pdf_path = build_pdf(tmp_path / "persian.pdf", [[persian_table]])
    outcome = convert_pdf(pdf_path, "opensource", make_settings())
    assert outcome.result.table_count == 1
    ws = _workbook(outcome).active
    values = {ws.cell(row=r, column=c).value for r in range(1, 4) for c in (1, 2)}
    assert "ردیف" in values
    assert "پرداخت" in values
    assert "بستانکار" in values


def test_e2e_debug_mode_saves_artifacts(make_settings, tmp_path):
    settings = make_settings(debug=True, debug_dir=tmp_path / "e2e_debug")
    pdf_path = build_pdf(tmp_path / "dbg.pdf", [[DEFAULT_TABLE]])
    outcome = convert_pdf(pdf_path, "opensource", settings)
    run_dirs = list((tmp_path / "e2e_debug").iterdir())
    assert run_dirs, "debug run directory not created"
    names = {p.name for p in run_dirs[0].iterdir()}
    assert "page_001.png" in names
    assert "tables_page_001.json" in names
    assert "structured_result.json" in names
    assert "result.xlsx" in names
    assert outcome.debug_dir is not None


def test_e2e_openai_engine_with_fake_client(make_settings, tmp_path):
    """Full pipeline with the OpenAI-compatible engine (offline fake)."""
    pdf_path = build_pdf(tmp_path / "oa.pdf", [[DEFAULT_TABLE]])
    settings = make_settings(openai_api_key="sk-test", openai_base_url=None)
    engine = OpenAICompatibleEngine(settings, client=make_client([JSON_TEXT]))
    outcome = convert_pdf(
        pdf_path, "openai", settings, engine=engine, filename="oa.pdf"
    )
    assert outcome.result.processing_mode == "openai"
    assert outcome.result.model_name == settings.openai_model
    assert outcome.result.table_count == 1
    assert outcome.validation.ok
    _, table = next(outcome.result.iter_tables())
    assert table.text_direction == "rtl"
    assert [c.header for c in table.columns] == ["ردیف", "شرح"]
    wb = _workbook(outcome)
    ws = wb.active
    assert ws.sheet_view.rightToLeft is True
    assert ws.cell(row=1, column=1).value == "ردیف"
    assert ws.cell(row=2, column=2).value == "پرداخت"


def test_e2e_same_normalized_shape_for_both_engines(make_settings, tmp_path):
    """Both engines emit the same normalized representation."""
    pdf_path = build_pdf(tmp_path / "shape.pdf", [[DEFAULT_TABLE]])
    open_outcome = convert_pdf(pdf_path, "opensource", make_settings())

    settings = make_settings(openai_api_key="sk-test")
    engine = OpenAICompatibleEngine(settings, client=make_client([JSON_TEXT]))
    llm_outcome = convert_pdf(pdf_path, "openai", settings, engine=engine)

    assert set(open_outcome.result.model_dump().keys()) == \
        set(llm_outcome.result.model_dump().keys())
    assert set(open_outcome.result.pages[0].model_dump().keys()) == \
        set(llm_outcome.result.pages[0].model_dump().keys())
    open_table = next(open_outcome.result.iter_tables())[1]
    llm_table = next(llm_outcome.result.iter_tables())[1]
    assert set(open_table.model_dump().keys()) == set(llm_table.model_dump().keys())
    assert open_outcome.result.processing_mode == "opensource"
    assert llm_outcome.result.processing_mode == "openai"

