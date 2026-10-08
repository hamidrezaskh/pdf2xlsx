"""Tests for XLSX generation (engine-independent layer)."""

from io import BytesIO

from openpyxl import load_workbook

from pdf2xlsx.schemas import (
    Cell,
    ColumnData,
    ExtractionResult,
    MergedCell,
    PageData,
    RowData,
    TableData,
)
from pdf2xlsx.xlsx import generate_xlsx


def _cell(r, c, value, **kw):
    return Cell(row_index=r, column_index=c, value=value, **kw)


def _load(data: bytes):
    return load_workbook(filename=BytesIO(data))


def make_rich_result() -> ExtractionResult:
    """2 pages, 3 tables, dynamic widths, empty cells, RTL, merge."""
    table1 = TableData(
        table_index=0,
        page_number=1,
        columns=[ColumnData(index=i, header=h) for i, h in enumerate(
            ["ID", "Name", "Qty", "Price"])],
        rows=[
            RowData(index=0, is_header=True, cells=[
                _cell(0, 0, "ID"), _cell(0, 1, "Name"), _cell(0, 2, "Qty"),
                _cell(0, 3, "Price")]),
            RowData(index=1, cells=[
                _cell(1, 0, "1"), _cell(1, 1, "Item A"), _cell(1, 2, "10"),
                _cell(1, 3, "4.50")]),
            RowData(index=2, cells=[
                _cell(2, 0, "2"), _cell(2, 1, ""), _cell(2, 2, "3"),
                _cell(2, 3, "12.00")]),
        ],
        text_direction="ltr",
    )
    table2 = TableData(  # narrow, RTL, Persian, different column count
        table_index=0,
        page_number=2,
        columns=[ColumnData(index=0, header="ردیف"), ColumnData(index=1, header="شرح")],
        rows=[
            RowData(index=0, is_header=True, cells=[
                _cell(0, 0, "ردیف"), _cell(0, 1, "شرح")]),
            RowData(index=1, cells=[_cell(1, 0, "۱"), _cell(1, 1, "پرداخت آب")]),
            RowData(index=2, cells=[_cell(2, 0, "۲"), _cell(2, 1, "")]),
        ],
        text_direction="rtl",
    )
    table3 = TableData(  # merged cells
        table_index=1,
        page_number=2,
        columns=[ColumnData(index=0, header="Region"), ColumnData(index=1, header="Value")],
        rows=[
            RowData(index=0, is_header=True, cells=[
                _cell(0, 0, "Region"), _cell(0, 1, "Value")]),
            RowData(index=1, cells=[
                _cell(1, 0, "Merged", colspan=2),
                _cell(1, 1, "covered", colspan=1),
            ]),
            RowData(index=2, cells=[_cell(2, 0, "a"), _cell(2, 1, "1")]),
        ],
        merged_cells=[
            MergedCell(row_index=1, column_index=0, rowspan=1, colspan=2, value="Merged")
        ],
    )
    return ExtractionResult(
        source_filename="report.pdf",
        page_count=2,
        processing_mode="opensource",
        pages=[
            PageData(page_number=1, tables=[table1]),
            PageData(page_number=2, tables=[table2, table3]),
        ],
    )


def test_sheet_per_table():
    wb = _load(generate_xlsx(make_rich_result()))
    assert wb.sheetnames == ["P1T1", "P2T1", "P2T2"]


def test_values_empty_cells_and_unicode():
    wb = _load(generate_xlsx(make_rich_result()))
    ws = wb["P1T1"]
    assert ws.cell(row=1, column=1).value == "ID"
    assert ws.cell(row=2, column=2).value == "Item A"
    assert ws.cell(row=3, column=2).value is None  # empty cell stays empty
    assert ws.cell(row=3, column=4).value == "12.00"

    rtl_ws = wb["P2T1"]
    assert rtl_ws.sheet_view.rightToLeft is True
    assert rtl_ws.cell(row=2, column=2).value == "پرداخت آب"
    assert rtl_ws.cell(row=3, column=2).value is None  # empty cell preserved


def test_header_styling_and_freeze():
    wb = _load(generate_xlsx(make_rich_result()))
    ws = wb["P1T1"]
    assert ws.cell(row=1, column=1).font.bold is True
    assert ws.freeze_panes == "A2"


def test_merged_cells_applied():
    wb = _load(generate_xlsx(make_rich_result()))
    ws = wb["P2T2"]
    merged = [str(r) for r in ws.merged_cells.ranges]
    # header row occupies sheet row 1, merged row_index=1 -> sheet row 2
    assert any(r.startswith("A2:B2") for r in merged)
    assert ws.cell(row=2, column=1).value == "Merged"
    assert ws.cell(row=2, column=2).value is None  # covered cell blank
    assert ws.cell(row=3, column=1).value == "a"  # next row intact


def test_dynamic_columns_per_sheet():
    wb = _load(generate_xlsx(make_rich_result()))
    assert wb["P1T1"].max_column == 4
    assert wb["P2T1"].max_column == 2


def test_column_widths_set():
    wb = _load(generate_xlsx(make_rich_result()))
    ws = wb["P1T1"]
    for letter in ("A", "B", "C", "D"):
        width = ws.column_dimensions[letter].width
        assert width and width >= 8


def test_headerless_table_writes_data_only():
    table = TableData(
        table_index=0,
        page_number=1,
        columns=[ColumnData(index=0), ColumnData(index=1)],
        rows=[RowData(index=0, cells=[_cell(0, 0, "a"), _cell(0, 1, "b")])],
    )
    result = ExtractionResult(
        source_filename="x.pdf", page_count=1,
        pages=[PageData(page_number=1, tables=[table])],
    )
    ws = _load(generate_xlsx(result)).active
    assert ws.cell(row=1, column=1).value == "a"
    assert ws.cell(row=1, column=2).value == "b"


def test_empty_result_produces_valid_workbook():
    result = ExtractionResult(source_filename="x.pdf", page_count=1, pages=[])
    wb = _load(generate_xlsx(result))
    assert len(wb.sheetnames) == 1
    assert wb.sheetnames[0] == "NoData"


def test_sheet_names_sanitized_and_unique():
    table = TableData(table_index=0, page_number=1,
                      columns=[ColumnData(index=0)],
                      rows=[RowData(index=0, cells=[_cell(0, 0, "v")])])
    other = TableData(table_index=1, page_number=1,
                      columns=[ColumnData(index=0)],
                      rows=[RowData(index=0, cells=[_cell(0, 0, "w")])])
    result = ExtractionResult(
        source_filename="x.pdf", page_count=1,
        pages=[PageData(page_number=1, tables=[table, other])],
    )
    names = _load(generate_xlsx(result)).sheetnames
    assert len(names) == len(set(n.lower() for n in names)) == 2
    assert all(len(n) <= 31 for n in names)


def test_no_engine_specific_logic_in_generator():
    """The generator consumes only the normalized model — no engine imports."""
    import inspect

    from pdf2xlsx.xlsx import generator

    source = inspect.getsource(generator)
    assert "pdf2xlsx.engines" not in source
    assert "from ..engines" not in source
    assert "google" not in source.lower()
    assert "easyocr" not in source.lower()
    assert "tesseract" not in source.lower()


