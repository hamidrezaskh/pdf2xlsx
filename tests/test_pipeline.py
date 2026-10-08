"""Tests for the orchestration pipeline."""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from pdf2xlsx.debug import DebugWriter
from pdf2xlsx.engines import ExtractionEngine
from pdf2xlsx.exceptions import ExtractionError, InvalidPdfError
from pdf2xlsx.pipeline import convert_pdf, xlsx_filename_for
from pdf2xlsx.schemas import (
    Cell,
    ColumnData,
    ExtractionResult,
    PageData,
    RowData,
    TableData,
)


class FakeEngine(ExtractionEngine):
    name = "fake"

    def __init__(self, settings, result_factory=None):
        super().__init__(settings)
        self.result_factory = result_factory
        self.calls = 0

    @property
    def model_name(self):
        return "fake-model"

    def extract(self, pdf, *, debug=None):
        self.calls += 1
        if self.result_factory:
            return self.result_factory(pdf)
        table = TableData(
            table_index=0,
            page_number=1,
            columns=[ColumnData(index=0, header="K"), ColumnData(index=1, header="V")],
            rows=[
                RowData(index=0, is_header=True, cells=[
                    Cell(row_index=0, column_index=0, value="K"),
                    Cell(row_index=0, column_index=1, value="V")]),
                RowData(index=1, cells=[
                    Cell(row_index=1, column_index=0, value="alpha"),
                    Cell(row_index=1, column_index=1, value="۱")]),
            ],
        )
        return ExtractionResult(
            source_filename=pdf.source_name,
            page_count=pdf.page_count,
            processing_mode=self.name,
            model_name=self.model_name,
            pages=[PageData(page_number=1, tables=[table])],
        )


def test_convert_pdf_end_to_end_with_injected_engine(grid_pdf, make_settings):
    steps = []
    outcome = convert_pdf(
        grid_pdf, "fake", make_settings(), engine=FakeEngine(make_settings()),
        on_progress=steps.append,
    )
    assert outcome.result.table_count == 1
    assert outcome.validation.ok
    assert outcome.duration_seconds >= 0
    assert outcome.xlsx_bytes[:2] == b"PK"  # zip/xlsx magic
    assert outcome.xlsx_filename == "grid.xlsx"
    joined = " | ".join(steps)
    assert "Processing mode" in joined
    assert "PDF loaded" in joined
    assert "Extraction finished" in joined
    assert "Validation finished" in joined
    assert "Generating XLSX" in joined
    wb = load_workbook(filename=BytesIO(outcome.xlsx_bytes))
    ws = wb.active
    assert ws.cell(row=2, column=1).value == "alpha"
    assert ws.cell(row=2, column=2).value == "۱"


def test_convert_pdf_invalid_bytes_raise(make_settings):
    with pytest.raises(InvalidPdfError):
        convert_pdf(b"definitely-not-a-pdf", "fake", make_settings(),
                    engine=FakeEngine(make_settings()))


def test_convert_pdf_unknown_mode_without_engine(grid_pdf, make_settings):
    from pdf2xlsx.exceptions import ConfigurationError

    with pytest.raises(ConfigurationError):
        convert_pdf(grid_pdf, "quantum-processor", make_settings())


def test_warnings_without_tables_surface_as_error(make_settings, tmp_path):
    from conftest import BorderlessSpec, build_pdf

    path = build_pdf(tmp_path / "emptyish.pdf", [[BorderlessSpec(
        rows=[["only"]], origins=[50.0], origin_y=50.0, line_pitch=20.0)]])

    def factory(pdf):
        return ExtractionResult(
            source_filename=pdf.source_name,
            page_count=pdf.page_count,
            processing_mode="fake",
            pages=[PageData(page_number=1, tables=[])],
            warnings=["Page 1 has no text layer and OCR is disabled; page skipped."],
        )

    engine = FakeEngine(make_settings(), result_factory=factory)
    with pytest.raises(ExtractionError) as exc:
        convert_pdf(path, "fake", make_settings(), engine=engine)
    assert "no text layer" in str(exc.value)


def test_no_tables_without_warnings_is_ok(make_settings, tmp_path):
    from conftest import TitleSpec, build_pdf

    only_title = build_pdf(tmp_path / "title.pdf", [[TitleSpec("Just a heading")]])

    def factory(pdf):
        return ExtractionResult(
            source_filename=pdf.source_name,
            page_count=pdf.page_count,
            processing_mode="fake",
            pages=[PageData(page_number=1, tables=[])],
        )

    outcome = convert_pdf(only_title, "fake", make_settings(),
                          engine=FakeEngine(make_settings(), result_factory=factory))
    assert outcome.result.table_count == 0
    assert outcome.xlsx_bytes[:2] == b"PK"


def test_debug_artifacts_saved(grid_pdf, make_settings, tmp_path):
    debug = DebugWriter(tmp_path / "dbg", run_id="pipeline_run")
    convert_pdf(grid_pdf, "fake", make_settings(),
                engine=FakeEngine(make_settings()), debug=debug)
    names = {p.name for p in debug.dir.iterdir()}
    assert "structured_result.json" in names
    assert "result.xlsx" in names


def test_debug_flag_creates_writer(grid_pdf, make_settings, tmp_path):
    settings = make_settings(debug=True, debug_dir=tmp_path / "auto_dbg")
    outcome = convert_pdf(grid_pdf, "fake", settings, engine=FakeEngine(settings))
    assert outcome.debug_dir is not None
    assert outcome.debug_dir.exists()
    assert (outcome.debug_dir / "result.xlsx").exists()


def test_xlsx_filename_sanitization():
    assert xlsx_filename_for("my report (final).pdf") == "my report (final).xlsx"
    assert xlsx_filename_for("a<b>:c.pdf") == "a_b__c.xlsx"
    assert xlsx_filename_for("") == "document.xlsx"
    assert xlsx_filename_for(None) == "document.xlsx"


def test_real_opensource_engine_via_pipeline(grid_pdf, make_settings):
    """Pipeline + registry create the opensource engine automatically."""
    outcome = convert_pdf(grid_pdf, "opensource", make_settings())
    assert outcome.result.processing_mode == "opensource"
    assert outcome.result.table_count == 1
    wb = load_workbook(filename=BytesIO(outcome.xlsx_bytes))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "ID"
    assert ws.cell(row=1, column=4).value == "Price"
    assert ws.cell(row=4, column=2).value is None  # empty cell round-trip

