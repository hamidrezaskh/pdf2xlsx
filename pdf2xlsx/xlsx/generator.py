"""XLSX generation from the normalized structured representation.

This layer is completely independent from the processing engine: it
consumes only ``ExtractionResult`` (Pydantic) and contains no
engine-specific logic.  Built on openpyxl.

Features: multiple tables/pages (one sheet per table), dynamic columns
and rows, empty cells, merged cells, RTL sheets, unicode text, header
styling and column auto-sizing.
"""

from __future__ import annotations

import logging
import re
from io import BytesIO
from typing import List, Set

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..exceptions import XlsxGenerationError
from ..schemas import ExtractionResult, TableData

logger = logging.getLogger(__name__)

_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
_MAX_SHEET_NAME = 31
_MAX_COL_WIDTH = 60


def _sheet_name(raw: str, used: Set[str]) -> str:
    name = _INVALID_SHEET_CHARS.sub("_", raw).strip() or "Sheet"
    name = name[:_MAX_SHEET_NAME]
    candidate = name
    counter = 2
    while candidate.lower() in used:
        suffix = f"_{counter}"
        candidate = name[: _MAX_SHEET_NAME - len(suffix)] + suffix
        counter += 1
    used.add(candidate.lower())
    return candidate


def _display_width(value: str) -> int:
    """Approximate display width (wide glyphs count double)."""
    width = 0
    for ch in value:
        width += 2 if ord(ch) > 0x0600 else 1  # rough RTL/CJK widening
    return width


def _has_header_row(table: TableData) -> bool:
    return any(row.is_header for row in table.rows) or any(col.header for col in table.columns)


def _write_table(ws, table: TableData, *, style: bool) -> None:
    n_cols = max(len(table.columns), max((len(r.cells) for r in table.rows), default=0))
    has_header = _has_header_row(table)
    header_in_rows = any(row.is_header for row in table.rows)
    # when the header row is part of the rows list it already occupies its
    # own sheet row; otherwise reserve row 1 for the column headers.
    offset = 1 if (has_header and not header_in_rows) else 0

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")

    if has_header:
        for col_index in range(n_cols):
            header = (
                table.columns[col_index].header
                if col_index < len(table.columns)
                else None
            )
            cell = ws.cell(row=1, column=col_index + 1, value=header or None)
            if style:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

    widths: List[int] = [8] * n_cols
    for row in table.rows:
        sheet_row = row.index + 1 + offset
        for cell in row.cells:
            value = cell.value if cell.value != "" else None
            ws_cell = ws.cell(row=sheet_row, column=cell.column_index + 1, value=value)
            if style and value is not None:
                ws_cell.alignment = Alignment(
                    vertical="center", wrap_text="\n" in str(value)
                )
            if value is not None:
                for part in str(value).split("\n"):
                    widths[cell.column_index] = max(
                        widths[cell.column_index], _display_width(part) + 2
                    )

    for col_index in range(n_cols):
        letter = get_column_letter(col_index + 1)
        ws.column_dimensions[letter].width = min(widths[col_index], _MAX_COL_WIDTH)

    # merged cells (values in anchor cell only; covered cells are blank)
    for merged in table.merged_cells:
        start_row = merged.row_index + 1 + offset
        start_col = merged.column_index + 1
        end_row = start_row + max(1, merged.rowspan) - 1
        end_col = start_col + max(1, merged.colspan) - 1
        if start_row < 1 or start_col < 1 or end_col > n_cols:
            continue
        try:
            ws.merge_cells(
                start_row=start_row,
                start_column=start_col,
                end_row=end_row,
                end_column=end_col,
            )
        except Exception as exc:  # pragma: no cover - defensive
            logger.warning("could not merge cells %s: %s", merged, exc)

    if has_header:
        ws.freeze_panes = "A2"


def generate_xlsx(result: ExtractionResult, *, style: bool = True) -> bytes:
    """Serialize a normalized extraction result into XLSX bytes."""
    try:
        workbook = Workbook()
        workbook.remove(workbook.active)
        used: Set[str] = set()
        written = 0

        for page in result.pages:
            for table in page.tables:
                if not table.rows and not table.columns:
                    logger.info(
                        "skipping empty table (page %d, index %d)",
                        page.page_number,
                        table.table_index,
                    )
                    continue
                name = _sheet_name(
                    f"P{page.page_number}T{table.table_index + 1}", used
                )
                ws = workbook.create_sheet(title=name)
                if table.text_direction == "rtl":
                    ws.sheet_view.rightToLeft = True
                _write_table(ws, table, style=style)
                written += 1

        if written == 0:
            workbook.create_sheet(title=_sheet_name("NoData", used))

        buffer = BytesIO()
        workbook.save(buffer)
        data = buffer.getvalue()
        logger.info("XLSX generated: %d sheet(s), %d byte(s)", written, len(data))
        return data
    except XlsxGenerationError:
        raise
    except Exception as exc:
        raise XlsxGenerationError(
            "The Excel file could not be generated from the extracted data.",
            detail=f"{type(exc).__name__}: {exc}",
        ) from exc


__all__ = ["generate_xlsx"]
